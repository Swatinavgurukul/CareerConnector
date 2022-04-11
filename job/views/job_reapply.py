from celery import app
from requests import api
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView
from resume.permissions import IsProfileOwnerOnly
from resume.models.profile import Profile
from resume.models.work import Work
from resume.models.education import Education
from resume.models.certificate import Certification
from resume.models.skill import Skills
from resume.models.achievement import Achievement
from resume.models.training import Training
from resume.models.language import Language
from core.helpers import api_response
from directsourcing.settings import MEDIA_ROOT
from fpdf import FPDF
from resume.views.pdf_file_insertion_creation import merge_pdfs, merge_docs
from core.models.users import User
from core.task import asynchronous_uploadResume
import docx
import datetime
from job.models.jobs import Job
from job.models.job_applications import JobApplication
from job.models.job_status import JobStatus
from bson.objectid import ObjectId
from resume.helpers import bimetric_score, recursive_lookup
from job.models.user_job_match import UserJobMatch
from recruiter.utils import get_real_sov_score
import pymongo

client = pymongo.MongoClient("mongodb://localhost:27017/")
# Database Name
db = client["dummydata"]
# Collection Name
col = db["resume"]
# Job Collection Name
job_col = db["job_parser"]
# Score Collection Name
score_col = db["score"]


def reapply_bimetricscore_update(job_id, user_id):
    try:
        user = User.objects.get(id=user_id)
        resume_id = user.mongo_object_id
        resume_data = col.find_one({"_id": ObjectId(resume_id)})
        parsed_resume = resume_data["data"]["Value"]["ResumeData"]

        job = Job.objects.get(id=job_id)
        job_id = job.mongo_object_id
        job_data = job_col.find_one({"_id": ObjectId(job_id)})
        parsed_job = job_data["data"]["Value"]["JobData"]

        data = bimetric_score(parsed_job, parsed_resume)
        # Inserting data into MongoDB
        resume_data = score_col.insert_one({"user_id": user.id, "job_id": job.id, "data": data})
        mongo_object_id = resume_data.inserted_id

        WeightedScore = recursive_lookup("WeightedScore", data)
        ReverseCompatibilityScore = recursive_lookup("ReverseCompatibilityScore", data)
        SovScore = recursive_lookup("SovScore", data)
        suggested_category_weights = recursive_lookup("SuggestedCategoryWeights", data)

        ScoreData = recursive_lookup("EnrichedScoreData", data)
        EnrichedScoreData = recursive_lookup("Taxonomies", ScoreData)
        if EnrichedScoreData is None:
            taxonomies_score = {"taxonomies_score": None}
        else:
            taxonomies_score = EnrichedScoreData.get("UnweightedScore")
            taxonomies_score = {"taxonomies_score": taxonomies_score}

        JobTitles = recursive_lookup("JobTitles", ScoreData)
        if JobTitles is None:
            jobtitle_score = {"jobtitle_score": None}
        else:
            jobtitle_score = JobTitles.get("UnweightedScore")
            jobtitle_score = {"jobtitle_score": jobtitle_score}

        skills = recursive_lookup("Skills", ScoreData)
        if skills is None:
            skill_score = {"skill_score": None}
        else:
            skill_score = skills.get("UnweightedScore")
            skill_score = {"skill_score": skill_score}

        management = recursive_lookup("ManagementLevel", ScoreData)
        if management is None:
            management_score = {"management_score": None}
        else:
            management_score = management.get("UnweightedScore")
            management_score = {"management_score": management_score}

        enrichedscore = {}
        enrichedscore.update(taxonomies_score)
        enrichedscore.update(jobtitle_score)
        enrichedscore.update(skill_score)
        enrichedscore.update(management_score)
        enriched_rcs_data = recursive_lookup("EnrichedRCSScoreData", data)
        certificate_score = recursive_lookup("Certifications", enriched_rcs_data)
        if certificate_score is None:
            certificate_rcs_score = {"certification": None}
        else:
            certificate_rcs_score = certificate_score.get("UnweightedScore")
            certificate_rcs_score = {"certification": certificate_rcs_score}

        education_score = recursive_lookup("Education", enriched_rcs_data)
        if education_score is None:
            education_rcs_score = {"education": None}
        else:
            education_rcs_score = education_score.get("UnweightedScore")
            education_rcs_score = {"education": education_rcs_score}

        enriched_rcs_score = recursive_lookup("Taxonomies", enriched_rcs_data)
        if enriched_rcs_score is None:
            taxonomies_rcs_score = {"industry": None}
        else:
            taxonomies_rcs_score = enriched_rcs_score.get("UnweightedScore")
            taxonomies_rcs_score = {"industry": taxonomies_rcs_score}

        title_rcs = recursive_lookup("JobTitles", enriched_rcs_data)
        if title_rcs is None:
            jobtitle_rcs_score = {"job_title": None}
        else:
            jobtitle_rcs_score = title_rcs.get("UnweightedScore")
            jobtitle_rcs_score = {"job_title": jobtitle_rcs_score}

        skill_rcs = recursive_lookup("Skills", enriched_rcs_data)
        if skill_rcs is None:
            skill_rcs_score = {"skills": None}
        else:
            skill_rcs_score = skill_rcs.get("UnweightedScore")
            skill_rcs_score = {"skills": skill_rcs_score}

        management_rcs = recursive_lookup("ManagementLevel", enriched_rcs_data)
        if management_rcs is None:
            management_rcs_score = {"management_level": None}
        else:
            management_rcs_score = management_rcs.get("UnweightedScore")
            management_rcs_score = {"management_level": management_rcs_score}

        rcs_score = {}
        rcs_score.update(certificate_rcs_score)
        rcs_score.update(education_rcs_score)
        rcs_score.update(taxonomies_rcs_score)
        rcs_score.update(jobtitle_rcs_score)
        rcs_score.update(skill_rcs_score)
        rcs_score.update(management_rcs_score)
        try:
            userjob_matchobj = UserJobMatch.objects.get(user=user, tenant=user.tenant, job=job)
            userjob_matchobj.weighted_score = WeightedScore
            userjob_matchobj.reverse_compatibility_score = (ReverseCompatibilityScore,)
            userjob_matchobj.sov_score = (SovScore,)
            userjob_matchobj.enriched_score = (enrichedscore,)
            userjob_matchobj.enriched_rcs_score = (rcs_score,)
            userjob_matchobj.unweighted_category_scores = (suggested_category_weights,)
            userjob_matchobj.suggested_score = (suggested_category_weights,)
            userjob_matchobj.mongo_object_id = (mongo_object_id,)
            userjob_matchobj.save()
        except UserJobMatch.DoesNotExist:
            pass
        user.custom_field3 = 1
        user.save()
    except Exception as e:
        return api_response(404, "Not found", str(e))


class JobReapplyView(APIView):
    permission_classes = (IsAuthenticated,)

    def put(self, request, slug, *args, **kwargs):
        profile_data_dict = {}
        user_id = request.user.id
        try:
            job = Job.objects.get(slug=slug)
        except Job.DoesNotExist:
            return api_response(404, "Job doesn't exist.", {})
        try:
            query_set = Profile.objects.get(user__id=user_id)
        except Profile.DoesNotExist:
            return api_response(404, "Page not found", {})
        work_obj = list(
            Work.objects.filter(user__id=user_id, user_created=True)
            .order_by("-start_date")
            .values(
                "title",
                "location__city",
                "city",
                "organization",
                "employment_type",
                "is_current",
                "description",
                "start_date",
                "end_date",
            )
        )
        for work in work_obj:
            for key in work.keys():
                if work[key] is None:
                    work[key] = ""

        edu_obj = list(
            Education.objects.filter(user_id=user_id, user_created=True)
            .order_by("-end_date")
            .values(
                "degree",
                "university_name",
                "university_type",
                "is_current",
                "start_date",
                "end_date",
            )
        )
        for edu in edu_obj:
            for key in edu.keys():
                if edu[key] is None:
                    edu[key] = ""
        certification_obj = list(
            Certification.objects.filter(user_id=user_id, user_created=True).values(
                "certification_name",
                "start_date",
                "end_date",
            )
        )
        for certification in certification_obj:
            for key in certification.keys():
                if certification[key] is None:
                    certification[key] = ""
        Skills_obj = list(Skills.objects.filter(user_id=user_id, user_created=True).values("skills"))
        for m in Skills_obj:
            if m["skills"] is None:
                m["skills"] = ""
        achievement_obj = list(
            Achievement.objects.filter(user_id=user_id, user_created=True).values(
                "achievement_name",
                "start_date",
                "end_date",
            )
        )
        for achievement in achievement_obj:
            for key in achievement.keys():
                if achievement[key] is None:
                    achievement[key] = ""
        training_obj = list(
            Training.objects.filter(user_id=user_id, user_created=True).values(
                "sov_training_name",
                "sov_training_description",
                "sov_training_startdate",
                "sov_training_enddate",
                "sov_training_type",
            )
        )
        for training in training_obj:
            for key in training.keys():
                if training[key] is None:
                    training[key] = ""
        language_obj = list(Language.objects.filter(user_id=user_id, user_created=True).values("name"))
        for q in language_obj:
            if q["name"] is None:
                q["name"] = ""
        # check if any extra user updated info is present
        if work_obj or edu_obj or Skills_obj or certification_obj or achievement_obj or training_obj or language_obj:
            profile_data_dict["EmploymentHistory"] = work_obj
            profile_data_dict["Education"] = edu_obj
            profile_data_dict["SkillsData"] = Skills_obj
            profile_data_dict["Certifications"] = certification_obj
            profile_data_dict["Achievements"] = achievement_obj
            profile_data_dict["Training"] = training_obj
            profile_data_dict["LanguageCompetencies"] = language_obj
            text_file_path = str(MEDIA_ROOT) + "edited.txt"

            with open(text_file_path, "w+", encoding="utf-8") as f:
                for key, val in profile_data_dict.items():
                    text = ""
                    for li in val:
                        for key1, val1 in li.items():
                            text = str(val1) + "," + "\n\t\t" + text
                    if text != "":
                        final_txt = "more" + str(key) + ": " + text
                        f.write(final_txt)
                        f.write("\n")

            f.close()

            custom_user = query_set.user.custom_field4
            if custom_user is not None:
                user_resume = query_set.user.custom_field4
            else:
                user_resume = query_set.user.resume_file
            output_path = ""
            user_resume_path = str(MEDIA_ROOT) + str(user_resume)
            if str(user_resume_path).endswith(".pdf"):
                edited_pdf_path = str(MEDIA_ROOT) + "resume/" + "edited" + ".pdf"
                f_pdf = FPDF()
                f_pdf.add_page()
                f_pdf.set_font("Times", size=12)
                try:
                    foo = open(text_file_path, "r+")
                    for x in foo:
                        f_pdf.cell(50, 5, txt=x.encode("utf-8").decode("ascii", "ignore"), ln=1, align="L")
                    f_pdf.output(edited_pdf_path)
                except UnicodeEncodeError:
                    pass
                finally:
                    foo.close()

                paths = [user_resume_path, edited_pdf_path]
                timestamp = datetime.datetime.now().timestamp()
                int_time = str(int(timestamp))
                name = str(query_set.user.resume_file).split(".pdf")
                output_path = str(MEDIA_ROOT) + name[0] + "_" + int_time + "_updated" + ".pdf"
                merge_pdfs(paths, output_path)

            if str(user_resume_path).endswith(".docx"):
                edited_docs_path = str(MEDIA_ROOT) + "resume/" + "edited" + ".docx"
                my_doc = docx.Document()
                for key, val in profile_data_dict.items():
                    text = ""
                    for li in val:
                        for key1, val1 in li.items():
                            text = str(val1) + "," + "\n\t\t" + text
                    if text != "":
                        final_txt = "more" + str(key) + ": " + text
                        my_doc.add_paragraph(final_txt)
                my_doc.save(edited_docs_path)
                paths = [user_resume_path, edited_docs_path]
                timestamp = datetime.datetime.now().timestamp()
                int_time = str(int(timestamp))
                name = str(query_set.user.resume_file).split(".docx")
                output_path = str(MEDIA_ROOT) + name[0] + "_" + int_time + ".docx"
                merge_docs(paths, output_path)
            l = output_path.split("resume")

            try:

                u_ob = User.objects.get(id=query_set.user.id)
                if str(user_resume).endswith(".pdf"):
                    u_ob.custom_field4 = "resume" + l[1]
                    u_ob.save()
                if str(user_resume).endswith(".docx"):
                    u_ob.custom_field4 = "resume" + l[1]
                    u_ob.save()

                asynchronous_uploadResume.delay(u_ob.id)

                u_ob.custom_field3 = 0
                u_ob.save()
            except User.DoesNotExist:
                return api_response(Response.status_code, "data", {"pdf is not created!!"})
            Work_Obj = list(
                Work.objects.filter(user__id=user_id, user_created=True)
                .order_by("-start_date")
                .values("id", "user__id")
            )
            for work in Work_Obj:
                wob = Work.objects.get(id=work["id"], user__id=work["user_id"], user_created=True)
                wob.user_created = False
                wob.save()
            Edu_Obj = list(
                Education.objects.filter(user_id=user_id, user_created=True)
                .order_by("-end_date")
                .values("id", "user_id")
            )
            for edu in Edu_Obj:
                eduob = Education.objects.get(id=edu["id"], user__id=edu["user_id"], user_created=True)
                eduob.user_created = False
                eduob.save()
            Certification_Obj = list(
                Certification.objects.filter(user_id=user_id, user_created=True).values("id", "user_id")
            )
            for certi in Certification_Obj:
                certiob = Certification.objects.get(id=certi["id"], user__id=certi["user_id"], user_created=True)
                certiob.user_created = False
                certiob.save()
            Skills_Obj = list(Skills.objects.filter(user_id=user_id, user_created=True).values("id", "user_id"))
            for skill in Skills_Obj:
                skillob = Skills.objects.get(id=skill["id"], user__id=skill["user_id"], user_created=True)
                skillob.user_created = False
                skillob.save()
            Achievement_Obj = list(
                Achievement.objects.filter(user_id=user_id, user_created=True).values("id", "user_id")
            )
            for achive in Achievement_Obj:
                achiveob = Achievement.objects.get(id=achive["id"], user__id=achive["user_id"], user_created=True)
                achiveob.user_created = False
                achiveob.save()
            Training_Obj = list(Training.objects.filter(user_id=user_id, user_created=True).values("id", "user_id"))
            for train in Training_Obj:
                training = Training.objects.get(id=train["id"], user__id=train["user_id"], user_created=True)
                training.user_created = False
                training.save()
            Language_Obj = list(Language.objects.filter(user_id=user_id, user_created=True).values("id", "user_id"))
            for lob in Language_Obj:
                lang = Language.objects.get(id=lob["id"], user__id=lob["user_id"], user_created=True)
                lang.user_created = False
                lang.save()

            reapply_bimetricscore_update(job.id, user_id)
            try:
                application = JobApplication.objects.get(tenant=request.user.tenant, user=request.user, job=job)
            except JobApplication.DoesNotExist:
                return api_response(404, "Not found", {})
            application.current_status = "applied"
            application.sov_data = get_real_sov_score(request.user.tenant, request.user, job)
            application.sov_score = application.sov_data["rcs"]
            application.weighted_score = application.sov_data["ws"]
            application.suggested_score = application.sov_data["suggested_score"]
            application.save()
            JobStatus.objects.create(
                user=request.user,
                job_tenant=job.tenant.id,
                tenant=request.user.tenant,
                job=job,
                application=application,
                status=application.current_status,
            )
            return api_response(200, "Successfully applied for job.", {})
        else:
            try:
                application = JobApplication.objects.get(
                    tenant=request.user.tenant,
                    user=request.user,
                    job__slug=slug,
                )
            except JobApplication.DoesNotExist:
                return api_response(404, "Not found", {})
            application.current_status = "applied"
            application.save()
            JobStatus.objects.create(
                user=request.user,
                job_tenant=job.tenant.id,
                tenant=request.user.tenant,
                job=job,
                application=application,
                status=application.current_status,
            )
            return api_response(200, "Successfully applied for job.", {})
